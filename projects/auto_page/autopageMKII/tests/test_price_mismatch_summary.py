import os
import sys
import shutil
import tempfile
import unittest
from unittest.mock import MagicMock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
import pandas as pd
from functions.accel_mode import AccelMode


class TestPriceMismatchSummary(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()
        self.test_excel = os.path.join(self.test_dir, "test_accel.xlsx")

        # Create basic Excel file with Sheet1
        df_main = pd.DataFrame({
            "หมายเลขคำสั่งซื้อ": ["ORD101", "ORD102", "ORD103"],
            "Tracking": ["TH111", "TH222", "TH333"]
        })
        with pd.ExcelWriter(self.test_excel, engine="openpyxl") as writer:
            df_main.to_excel(writer, sheet_name="Sheet1", index=False)

        mock_app = MagicMock()
        mock_app.tracking_from_data = ["TH111"]
        mock_app.items = []
        mock_app.financials = None

        self.accel = AccelMode(mock_app)
        self.accel.accel_file_dir = self.test_excel
        self.accel.main_sheet_name = "Sheet1"

    def tearDown(self):
        if os.path.exists(self.test_dir):
            shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_price_mismatch_summary_creation(self):
        """ทดสอบว่าเมื่อเกิด PRICE_MISMATCH จะสร้างชีต Price_Mismatch_Summary พร้อมสูตร COUNTIFS ถูกต้อง"""
        self.accel.record_failed_order(
            order="ORD101",
            reason="ขอวิธีปรับราคาครับ (มี SKU ใน CP_data แล้ว แต่ยังไม่มี CP/ส่วนลดให้เลือก)",
            category="PRICE_MISMATCH",
            sku_pattern="MNL-002387",
            price=500.0
        )

        wb = openpyxl.load_workbook(self.test_excel)
        self.assertIn("Price_Mismatch_Summary", wb.sheetnames)
        self.assertIn("Processed_Logs", wb.sheetnames)

        ws_sum = wb["Price_Mismatch_Summary"]
        headers = [ws_sum.cell(row=1, column=c).value for c in range(1, 4)]
        self.assertEqual(headers, ['sku_pattern', 'expected_price', 'remaining_failed_count'])

        # ตรวจสอบข้อมูลแถวที่ 2
        sku_val = ws_sum.cell(row=2, column=1).value
        price_val = ws_sum.cell(row=2, column=2).value
        formula_val = ws_sum.cell(row=2, column=3).value

        self.assertEqual(sku_val, "MNL-002387")
        self.assertEqual(price_val, "500")
        expected_formula = (
            '=COUNTIFS(Processed_Logs!G:G, "Failed", '
            'Processed_Logs!H:H, "PRICE_MISMATCH", '
            'Processed_Logs!I:I, "*[SKU: " & A2 & "]*", '
            'Processed_Logs!E:E, B2)'
        )
        self.assertEqual(formula_val, expected_formula)

        # ตรวจสอบ Processed_Logs
        ws_proc = wb["Processed_Logs"]
        # Col E (5) = price, Col G (7) = status, Col H (8) = error_category, Col I (9) = remark
        proc_price = ws_proc.cell(row=2, column=5).value
        proc_status = ws_proc.cell(row=2, column=7).value
        proc_cat = ws_proc.cell(row=2, column=8).value
        proc_remark = ws_proc.cell(row=2, column=9).value

        self.assertEqual(str(proc_price), "500")
        self.assertEqual(proc_status, "Failed")
        self.assertEqual(proc_cat, "PRICE_MISMATCH")
        self.assertIn("[SKU: MNL-002387]", proc_remark)

        wb.close()

    def test_combo_ink_set_price_mismatch(self):
        """ทดสอบชุดเซ็ตหมึก เช่น SP2-001845+SP2-001846+SP2-001847+SP2-001848"""
        combo_sku = "SP2-001845+SP2-001846+SP2-001847+SP2-001848"
        self.accel.record_failed_order(
            order="ORD102",
            reason="ราคาไม่ตรงเป้าหมาย",
            category="PRICE_MISMATCH",
            sku_pattern=combo_sku,
            price=945.0
        )

        wb = openpyxl.load_workbook(self.test_excel)
        ws_sum = wb["Price_Mismatch_Summary"]
        self.assertEqual(ws_sum.cell(row=2, column=1).value, combo_sku)
        self.assertEqual(ws_sum.cell(row=2, column=2).value, "945")

        ws_proc = wb["Processed_Logs"]
        self.assertIn(f"[SKU: {combo_sku}]", ws_proc.cell(row=2, column=9).value)
        wb.close()

    def test_duplicate_price_mismatch_does_not_duplicate_summary_rows(self):
        """ทดสอบว่าเมื่อเกิด PRICE_MISMATCH ใน SKU Pattern และราคาเดียวกันซ้ำจากหลายออเดอร์ จะไม่สร้างแถวซ้ำใน Summary"""
        self.accel.record_failed_order(
            order="ORD_A",
            reason="ราคาไม่ตรง",
            category="PRICE_MISMATCH",
            sku_pattern="SKU_SET_A",
            price=1200.0
        )

        # Reset debounce tracker to simulate another order
        self.accel._last_failed_record = (None, 0)

        self.accel.record_failed_order(
            order="ORD_B",
            reason="ราคาไม่ตรง",
            category="PRICE_MISMATCH",
            sku_pattern="SKU_SET_A",
            price=1200.0
        )

        wb = openpyxl.load_workbook(self.test_excel)
        ws_sum = wb["Price_Mismatch_Summary"]
        # แถว 1 คือ Header, แถว 2 คือ SKU_SET_A -> ต้องไม่มีแถว 3
        self.assertEqual(ws_sum.max_row, 2)

        # แต่ Processed_Logs ต้องมีทั้ง 2 ออเดอร์
        ws_proc = wb["Processed_Logs"]
        self.assertEqual(ws_proc.max_row, 3)
        wb.close()

    def test_auto_extraction_from_main_app_items_excluding_shipping(self):
        """ทดสอบการดึง SKU Pattern และ expected_price อัตโนมัติจาก main_app.items พร้อมตัดค่าส่ง SV0-000101 ออก"""
        mock_items = [
            {
                "เลขอ้างอิง SKU (SKU Reference No.)": "SP2-001845+SP2-001846+SP2-001847+SP2-001848",
                "ราคาขายสุทธิ": "945",
                "จำนวน": "1",
                "ส่วนลดจาก Shopee": "0"
            },
            {
                "เลขอ้างอิง SKU (SKU Reference No.)": "SV0-000101",
                "ราคาขายสุทธิ": "40",
                "จำนวน": "1",
                "ส่วนลดจาก Shopee": "0"
            }
        ]
        self.accel.main_app.items = mock_items

        self.accel.record_failed_order(
            order="ORD_AUTO_1",
            reason="ขอวิธีปรับราคาครับ (มี SKU ใน CP_data แล้ว แต่ยังไม่มี CP/ส่วนลดให้เลือก)"
        )

        wb = openpyxl.load_workbook(self.test_excel)
        ws_sum = wb["Price_Mismatch_Summary"]
        self.assertEqual(ws_sum.max_row, 2)
        self.assertEqual(ws_sum.cell(row=2, column=1).value, "SP2-001845+SP2-001846+SP2-001847+SP2-001848")
        self.assertEqual(ws_sum.cell(row=2, column=2).value, "945")

        # ตรวจสอบว่าไม่มี SV0-000101 บันทึกลงใน Price_Mismatch_Summary
        all_summary_skus = [ws_sum.cell(row=r, column=1).value for r in range(2, ws_sum.max_row + 1)]
        self.assertNotIn("SV0-000101", all_summary_skus)
        wb.close()

    def test_retry_success_updates_processed_logs_status(self):
        """ทดสอบว่าเมื่อ retry จนออกบิลสำเร็จ Processed_Logs จะเปลี่ยนสถานะเป็น Completed ทำให้สูตร COUNTIFS คัดยอดออก"""
        # Step 1: Failed attempt
        self.accel.record_failed_order(
            order="ORD_RETRY_99",
            reason="ขอวิธีปรับราคาครับ",
            category="PRICE_MISMATCH",
            sku_pattern="MNL-999",
            price=750.0
        )

        df_proc1 = pd.read_excel(self.test_excel, sheet_name="Processed_Logs")
        self.assertEqual(len(df_proc1), 1)
        self.assertEqual(df_proc1.iloc[0]["status"], "Failed")
        self.assertEqual(df_proc1.iloc[0]["error_category"], "PRICE_MISMATCH")

        # Step 2: Retry completed
        self.accel.record_completed_order(
            order="ORD_RETRY_99",
            tracking="TH_SUCCESS",
            bill_no="B01-999",
            status="Completed",
            price="750",
            serials="SN_OK_1"
        )

        df_proc2 = pd.read_excel(self.test_excel, sheet_name="Processed_Logs")
        self.assertEqual(len(df_proc2), 1)
        self.assertEqual(df_proc2.iloc[0]["status"], "Completed")
        self.assertEqual(str(df_proc2.iloc[0]["error_category"]).replace("nan", ""), "")

        # Summary sheet still exists and keeps the formula intact
        wb = openpyxl.load_workbook(self.test_excel)
        ws_sum = wb["Price_Mismatch_Summary"]
        self.assertEqual(ws_sum.cell(row=2, column=1).value, "MNL-999")
        self.assertIn("COUNTIFS", str(ws_sum.cell(row=2, column=3).value))
        wb.close()

    def test_valid_sheets_excludes_price_mismatch_summary(self):
        """ทดสอบว่า _get_main_sheet_name จะไม่เลือก Price_Mismatch_Summary เป็นชีตออเดอร์หลัก"""
        wb = openpyxl.load_workbook(self.test_excel)
        wb.create_sheet("Price_Mismatch_Summary")
        wb.save(self.test_excel)
        wb.close()

        main_sheet = self.accel._get_main_sheet_name(self.test_excel)
        self.assertEqual(main_sheet, "Sheet1")


if __name__ == "__main__":
    unittest.main()
